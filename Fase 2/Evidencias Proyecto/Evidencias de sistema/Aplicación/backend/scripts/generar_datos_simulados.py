"""Genera los informes de deudores simulados de la TGR para desarrollar y probar TTDH Automation.

La TGR envía por correo un informe con los deudores (PDF, escaneo, foto, Excel, Word, CSV…). Este script
crea informes ficticios en todos esos formatos, con casos válidos y casos de borde, una base simulada de
cuentas de Banco Genérico y un manifiesto con el resultado esperado de cada archivo y deudor, que sirve
como referencia para las pruebas de la extracción y del ETL.

Columnas del informe: las que, según el equipo, vienen de la TGR en el CSV de Tomás (RUT, nombre, ID de
demanda, N° de resolución, monto solicitado, tipo de embargo, expediente, tesorería y fecha del oficio).

Uso, desde la carpeta backend/:
    python scripts/generar_datos_simulados.py                 # escribe en ../data/simulados
    python scripts/generar_datos_simulados.py --grande 12000  # agrega un informe de volumen

Con la misma semilla el resultado es siempre el mismo. Solo se sobrescriben los archivos que
genera este script; nunca se borra nada de la carpeta de salida.

Los RUT se generan entre 40.000.000 y 49.999.999, fuera de los rangos que hoy se asignan a
personas (bajo 30 millones) y empresas (desde 50 millones), para no asociar por accidente a
una persona real con una deuda.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import re
import shutil
import sys
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont

CARPETA_SALIDA = Path(__file__).resolve().parents[2] / "data" / "simulados"

COLUMNAS = (
    "rut", "nombre", "id_demanda", "n_resolucion", "monto_solicitado",
    "tipo_embargo", "expediente", "tesoreria", "fecha_oficio",
)
TITULOS = {
    "rut": "RUT", "nombre": "Nombre", "id_demanda": "ID Demanda", "n_resolucion": "N° Resolución",
    "monto_solicitado": "Monto Solicitado", "tipo_embargo": "Tipo Embargo", "expediente": "Expediente",
    "tesoreria": "Tesorería", "fecha_oficio": "Fecha Oficio",
}
TIPOS_EMBARGO = ("RETENCIÓN", "EMBARGO")
# (tesorería, localidad del expediente)
TESORERIAS = (
    ("TESORERÍA REGIONAL DE SANTIAGO", "SANTIAGO"),
    ("TESORERÍA REGIONAL DE VALPARAÍSO", "VALPARAÍSO"),
    ("TESORERÍA REGIONAL DEL BIOBÍO", "CONCEPCIÓN"),
    ("TESORERÍA PROVINCIAL DE TALAGANTE", "TALAGANTE"),
    ("TESORERÍA PROVINCIAL DEL MAIPO", "SAN BERNARDO"),
    ("TESORERÍA REGIONAL DE ANTOFAGASTA", "ANTOFAGASTA"),
)
CODIGO_BANCO_GENERICO = "099"
FECHA_GENERACION = date(2026, 10, 1)
FECHA_METADATOS = datetime(2026, 10, 1, 9, 0)  # fija, para que los archivos sean reproducibles
TITULO_INFORME = "Tesorería General de la República - Informe de deudores"
AVISO = "DOCUMENTO SIMULADO - SIN VALIDEZ - SOLO PARA PRUEBAS"

# Resultado esperado de cada deudor después de la extracción y el ETL.
CARGADO, RECHAZADO, DUPLICADO, REVISION = "Cargado", "Rechazado", "Duplicado", "Revisión manual"

NOMBRES = (
    "María", "José", "Juan", "Camila", "Francisca", "Benjamín", "Sofía", "Matías", "Valentina",
    "Tomás", "Javiera", "Agustín", "Constanza", "Vicente", "Antonia", "Joaquín", "Catalina",
    "Ignacio", "Fernanda", "Sebastián",
)
APELLIDOS = (
    "González", "Muñoz", "Rojas", "Díaz", "Pérez", "Soto", "Contreras", "Silva", "Martínez",
    "Sepúlveda", "Morales", "Rodríguez", "López", "Fuentes", "Hernández", "Torres", "Araya",
    "Flores", "Espinoza", "Valenzuela", "Castillo", "Núñez", "Ibáñez", "Peña",
)


def digito_verificador(numero: int) -> str:
    """Dígito verificador del RUT según el algoritmo Módulo 11 (RN01)."""
    suma, factor = 0, 2
    for digito in reversed(str(numero)):
        suma += int(digito) * factor
        factor = 2 if factor == 7 else factor + 1
    resto = 11 - suma % 11
    return {11: "0", 10: "K"}.get(resto, str(resto))


def rut_con_puntos(rut: str) -> str:
    numero, dv = rut.split("-")
    return f"{int(numero):,}".replace(",", ".") + f"-{dv}"


def formato_clp(monto: int) -> str:
    return "$ " + f"{monto:,}".replace(",", ".")


@dataclass
class Registro:
    """Un deudor del informe, con sus valores normalizados y el resultado esperado."""

    rut: str
    nombre: str
    id_demanda: str
    n_resolucion: str
    monto: int | None
    tipo_embargo: str
    expediente: str
    tesoreria: str
    fecha_oficio: date | None
    # Cómo aparece un campo en el archivo cuando difiere del valor normalizado (casos de borde).
    escrito: dict[str, str] = field(default_factory=dict)
    estado: str = CARGADO
    motivo: str | None = None
    caso: str | None = None

    def normalizado(self) -> dict[str, object]:
        return {
            "rut": self.rut, "nombre": self.nombre, "id_demanda": self.id_demanda,
            "n_resolucion": self.n_resolucion, "monto_solicitado": self.monto, "tipo_embargo": self.tipo_embargo,
            "expediente": self.expediente, "tesoreria": self.tesoreria,
            "fecha_oficio": self.fecha_oficio.isoformat() if self.fecha_oficio else None,
        }

    def texto(self, columna: str) -> str:
        """Valor tal como se escribe en los archivos de datos (CSV, TXT, Excel)."""
        if columna in self.escrito:
            return self.escrito[columna]
        if columna == "fecha_oficio":
            return self.fecha_oficio.strftime("%d-%m-%Y") if self.fecha_oficio else ""
        valor = self.normalizado()[columna]
        return "" if valor is None else str(valor)

    def texto_impreso(self, columna: str) -> str:
        """Valor tal como se imprime en los PDF, imágenes y Word (formato chileno)."""
        if columna in self.escrito:
            return self.escrito[columna]
        if columna == "rut" and self.rut:
            return rut_con_puntos(self.rut)
        if columna == "monto_solicitado" and self.monto is not None:
            return formato_clp(self.monto)
        return self.texto(columna)

    def para_manifiesto(self) -> dict[str, object]:
        fila = {**self.normalizado(), "estado_esperado": self.estado}
        if self.motivo:
            fila["motivo"] = self.motivo
        if self.caso:
            fila["caso"] = self.caso
        if self.escrito:
            fila["valor_en_archivo"] = dict(self.escrito)
        return fila


def copia(registro: Registro, **cambios: object) -> Registro:
    return replace(registro, escrito=dict(registro.escrito), **cambios)


@dataclass
class Informe:
    lote: str
    registros: list[Registro]
    cantidad_declarada: int | None = None  # None: la cantidad real del detalle
    suma_declarada: int | None = None  # None: la suma real de los montos del detalle

    @property
    def cantidad(self) -> int:
        return len(self.registros) if self.cantidad_declarada is None else self.cantidad_declarada

    @property
    def suma(self) -> int:
        if self.suma_declarada is not None:
            return self.suma_declarada
        return sum(r.monto for r in self.registros if r.monto is not None)

    def lineas_encabezado(self) -> list[list[str]]:
        return [[TITULO_INFORME], [AVISO],
                [f"Lote: {self.lote}", f"Fecha de generación: {FECHA_GENERACION:%d-%m-%Y}",
                 f"Registros: {self.cantidad}"]]

    def linea_totales(self) -> list[str]:
        return ["TOTALES", f"Registros: {self.cantidad}", f"Monto total: {formato_clp(self.suma)}"]


class Generador:
    """Produce deudores ficticios reproducibles a partir de una semilla."""

    def __init__(self, semilla: int) -> None:
        self.rng = random.Random(semilla)
        self._ultima_demanda = 480_000
        self._ultimo_lote = 0

    def lote(self) -> str:
        self._ultimo_lote += 1
        return f"TGR-{FECHA_GENERACION:%Y%m%d}-{self._ultimo_lote:02d}"

    def rut(self, dv: str | None = None) -> str:
        while True:
            numero = self.rng.randint(40_000_000, 49_999_999)
            digito = digito_verificador(numero)
            if dv is None or digito == dv:
                return f"{numero}-{digito}"

    def registro(self) -> Registro:
        self._ultima_demanda += self.rng.randint(1, 40)
        tesoreria, localidad = self.rng.choice(TESORERIAS)
        return Registro(
            rut=self.rut(),
            nombre=f"{self.rng.choice(NOMBRES)} {self.rng.choice(APELLIDOS)} {self.rng.choice(APELLIDOS)}",
            id_demanda=str(self._ultima_demanda),
            n_resolucion=str(self.rng.randint(10_000, 99_999)),
            monto=self.rng.randint(20, 3_000) * 1_000,
            tipo_embargo=self.rng.choice(TIPOS_EMBARGO),
            expediente=f"{self.rng.randint(100, 99_999)}-{self.rng.randint(2019, 2026)} {localidad}",
            tesoreria=tesoreria,
            fecha_oficio=FECHA_GENERACION - timedelta(days=self.rng.randint(1, 30)),
        )

    def registros(self, cantidad: int) -> list[Registro]:
        return [self.registro() for _ in range(cantidad)]

    def informe(self, cantidad: int) -> Informe:
        return Informe(self.lote(), self.registros(cantidad))


def casos_de_borde(gen: Generador) -> list[Registro]:
    """Deudores válidos de base más un caso por cada regla de la extracción y del ETL."""
    registros = gen.registros(10)

    def agregar(caso: str, estado: str = CARGADO, motivo: str | None = None) -> Registro:
        registro = gen.registro()
        registro.caso, registro.estado, registro.motivo = caso, estado, motivo
        registros.append(registro)
        return registro

    r = agregar("RUT con puntos")
    r.escrito["rut"] = rut_con_puntos(r.rut)
    r = agregar("RUT con dígito verificador k minúscula")
    r.rut = gen.rut(dv="K")
    r.escrito["rut"] = r.rut.lower()
    r = agregar("RUT con espacios alrededor")
    r.escrito["rut"] = f"  {r.rut} "

    r = agregar("Dígito verificador incorrecto", RECHAZADO, "RUT o DV inválido")
    numero, dv = r.rut.split("-")
    r.rut = f"{numero}-{'1' if dv != '1' else '2'}"
    r = agregar("RUT con letras en el número", RECHAZADO, "RUT o DV inválido")
    r.rut = r.rut[:2] + "A" + r.rut[3:]
    r = agregar("RUT vacío", RECHAZADO, "RUT o DV inválido")
    r.rut = ""

    r = agregar("Sin monto", RECHAZADO, "Falta MONTO SOLICITADO (RN03)")
    r.monto = None
    r = agregar("Monto con formato de moneda")
    r.escrito["monto_solicitado"] = formato_clp(r.monto)
    r = agregar("Monto escrito en palabras", RECHAZADO, "MONTO SOLICITADO no es un número entero")
    r.monto = None
    r.escrito["monto_solicitado"] = "mil pesos"

    r = agregar("Fecha dd/mm/aaaa")
    r.escrito["fecha_oficio"] = r.fecha_oficio.strftime("%d/%m/%Y")
    r = agregar("Fecha inexistente", RECHAZADO, "FECHA OFICIO no es dd-mm-aaaa")
    r.fecha_oficio = None
    r.escrito["fecha_oficio"] = "31-02-2026"
    r = agregar("Sin fecha del oficio", REVISION, "RN05 - Falta FECHA OFICIO: se carga marcado para revisión")
    r.fecha_oficio = None

    r = agregar("Expediente sin formato número-año localidad", RECHAZADO, "EXPEDIENTE sin formato")
    r.expediente = "SIN ROL"
    r = agregar("Tesorería sin el formato habitual: se carga igual (sin tipo ni localidad)")
    r.tesoreria = "OFICINA CENTRAL"

    r = agregar("Nombre con ñ y tildes")
    r.nombre = "Begoña Ñancupil Ibáñez"
    r = agregar("Sin nombre", REVISION, "RN05 - Falta NOMBRE CLIENTE: se carga marcado para revisión")
    r.nombre = ""

    base = registros[1]
    registros.append(copia(
        base, expediente=f"{gen.rng.randint(100, 99_999)}-2025 SANTIAGO", id_demanda=str(int(base.id_demanda) + 7),
        caso=f"Mismo RUT que {base.rut} con otro expediente: dos deudas (MULTIPLES ROLES = SI)",
    ))
    original = registros[2]
    registros.append(copia(
        original, estado=DUPLICADO, caso="Repetición exacta de otro deudor del mismo informe",
        motivo=f"Mismo RUT, expediente y fecha de ingreso que {original.rut}",
    ))
    return registros


# --------------------------------------------------------------- escritura de archivos


def escribir_delimitado(
    ruta: Path, informe: Informe, *, delimitador: str = ";", encoding: str = "utf-8",
    con_totales: bool = True, columnas: tuple[str, ...] = COLUMNAS, titulos: dict[str, str] | None = None,
) -> None:
    titulos = {**TITULOS, **(titulos or {})}
    with ruta.open("w", encoding=encoding, newline="") as archivo:
        escritor = csv.writer(archivo, delimiter=delimitador)
        escritor.writerows(informe.lineas_encabezado())
        escritor.writerow([titulos[c] for c in columnas])
        escritor.writerows([r.texto(c) for c in columnas] for r in informe.registros)
        if con_totales:
            escritor.writerow(informe.linea_totales())


def escribir_excel(ruta: Path, informe: Informe) -> None:
    """Montos y fechas válidos como celdas numéricas y de fecha, igual que un Excel real."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Deudores"
    for linea in informe.lineas_encabezado():
        hoja.append(linea)
    hoja.append([TITULOS[c] for c in COLUMNAS])
    for registro in informe.registros:
        fila: list[object] = [registro.texto(c) for c in COLUMNAS]
        if registro.monto is not None and "monto_solicitado" not in registro.escrito:
            fila[COLUMNAS.index("monto_solicitado")] = registro.monto
        if registro.fecha_oficio is not None and "fecha_oficio" not in registro.escrito:
            fila[COLUMNAS.index("fecha_oficio")] = registro.fecha_oficio
        hoja.append(fila)
    hoja.append(informe.linea_totales())
    libro.properties.creator = "TTDH Automation - datos simulados"
    libro.properties.created = FECHA_METADATOS
    libro.save(ruta)
    _fijar_fechas_office(ruta)


def escribir_word(ruta: Path, informe: Informe) -> None:
    """Informe en Word: encabezado en párrafos, deudores en una tabla y totales al final."""
    documento = Document()
    estilo = documento.styles["Normal"]
    estilo.font.name, estilo.font.size = "Arial", Pt(8)
    documento.add_heading(TITULO_INFORME, level=1)
    documento.add_paragraph(AVISO)
    documento.add_paragraph("    ".join(informe.lineas_encabezado()[2]))
    tabla = documento.add_table(rows=1, cols=len(COLUMNAS))
    tabla.style = "Table Grid"
    for celda, columna in zip(tabla.rows[0].cells, COLUMNAS):
        celda.text = TITULOS[columna]
    for registro in informe.registros:
        for celda, columna in zip(tabla.add_row().cells, COLUMNAS):
            celda.text = registro.texto_impreso(columna)
    documento.add_paragraph("    ".join(informe.linea_totales()))
    documento.core_properties.author = "TTDH Automation - datos simulados"
    documento.core_properties.created = documento.core_properties.modified = FECHA_METADATOS
    documento.save(ruta)
    _fijar_fechas_office(ruta)


def _fijar_fechas_office(ruta: Path) -> None:
    """openpyxl y python-docx graban la hora actual al guardar (en los metadatos y en el zip). Se reemplaza
    por una fija para que regenerar los datos no cambie los archivos."""
    modificado = FECHA_METADATOS.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(ruta) as original:
        entradas = [(info.filename, original.read(info)) for info in original.infolist()]
    with zipfile.ZipFile(ruta, "w") as nuevo:
        for nombre, datos in entradas:
            if nombre == "docProps/core.xml":
                datos = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + modificado, datos)
            info = zipfile.ZipInfo(nombre, date_time=FECHA_METADATOS.timetuple()[:6])
            nuevo.writestr(info, datos, compress_type=zipfile.ZIP_DEFLATED)


ANCHOS_PDF_MM = (23, 50, 18, 21, 25, 21, 40, 56, 20)


def escribir_pdf_texto(ruta: Path, informe: Informe) -> None:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_creation_date(FECHA_METADATOS)
    pdf.set_title(f"Informe de deudores {informe.lote}")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 7, TITULO_INFORME, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 5, AVISO, new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, "    ".join(informe.lineas_encabezado()[2]), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 7)
    for columna, ancho in zip(COLUMNAS, ANCHOS_PDF_MM):
        pdf.cell(ancho, 6, TITULOS[columna], border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7)
    for registro in informe.registros:
        for columna, ancho in zip(COLUMNAS, ANCHOS_PDF_MM):
            alineacion = "R" if columna == "monto_solicitado" else "L"
            pdf.cell(ancho, 6, registro.texto_impreso(columna), border=1, align=alineacion)
        pdf.ln()
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 6, "    ".join(informe.linea_totales()))
    pdf.output(str(ruta))


def _fuente(tamano: int, negrita: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    nombres = ("arialbd.ttf", "DejaVuSans-Bold.ttf") if negrita else ("arial.ttf", "DejaVuSans.ttf")
    for nombre in nombres:
        try:
            return ImageFont.truetype(nombre, tamano)
        except OSError:
            continue
    return ImageFont.load_default(size=tamano)


# Deja margen a los lados: la perspectiva de la foto recorta los bordes de la hoja.
ANCHOS_IMAGEN_PX = (190, 360, 140, 160, 190, 160, 310, 460, 160)


def dibujar_informe(informe: Informe) -> tuple[Image.Image, list[tuple[int, int, int, int]]]:
    """Dibuja el informe como una hoja A4 apaisada a 200 dpi.

    Devuelve la imagen y el recuadro de la celda de monto de cada deudor.
    """
    imagen = Image.new("RGB", (2339, 1654), "white")
    lapiz = ImageDraw.Draw(imagen)
    x0, y = 100, 80
    lapiz.text((x0, y), TITULO_INFORME, font=_fuente(40, True), fill="black")
    y += 60
    lapiz.text((x0, y), AVISO, font=_fuente(24), fill="black")
    y += 40
    lapiz.text((x0, y), "     ".join(informe.lineas_encabezado()[2]), font=_fuente(28), fill="black")
    y += 70
    alto_fila = 44
    filas = [[TITULOS[c] for c in COLUMNAS]] + [[r.texto_impreso(c) for c in COLUMNAS] for r in informe.registros]
    cajas_monto = []
    for indice, valores in enumerate(filas):
        fuente = _fuente(22, negrita=indice == 0)
        x = x0
        for columna, ancho, valor in zip(COLUMNAS, ANCHOS_IMAGEN_PX, valores):
            caja = (x, y, x + ancho, y + alto_fila)
            lapiz.rectangle(caja, outline="black", width=2)
            lapiz.text((x + 8, y + 10), valor, font=fuente, fill="black")
            if columna == "monto_solicitado" and indice > 0:
                cajas_monto.append(caja)
            x += ancho
        y += alto_fila
    y += 30
    lapiz.text((x0, y), "     ".join(informe.linea_totales()), font=_fuente(28, True), fill="black")
    return imagen, cajas_monto


def escribir_pdf_escaneado(ruta: Path, informe: Informe, rng: random.Random) -> None:
    """Simula un escaneo: gris, levemente girado, con ruido y algo de desenfoque.

    El ruido sale de `rng` (Image.effect_noise no acepta semilla y haría el archivo distinto cada vez).
    """
    hoja, _ = dibujar_informe(informe)
    gris = hoja.convert("L").rotate(0.6, resample=Image.Resampling.BICUBIC, fillcolor=255)
    ruido = Image.frombytes("L", gris.size, rng.randbytes(gris.width * gris.height))
    gris = Image.blend(gris, ruido, 0.08)
    gris = gris.filter(ImageFilter.GaussianBlur(0.5))
    gris.save(ruta, "PDF", resolution=200.0, title=f"Informe escaneado {informe.lote}",
              creationDate=FECHA_METADATOS.timetuple(), modDate=FECHA_METADATOS.timetuple())


def escribir_fotografia(ruta: Path, informe: Informe, filas_manchadas: list[int]) -> None:
    """Simula una foto de celular: perspectiva, sombra, bajo contraste, desenfoque y manchas.

    Una mancha uniforme no basta (el OCR la separa como si fuera fondo), así que el monto
    manchado además se corre: con ese desenfoque Tesseract deja de leerlo con confianza.
    """
    hoja, cajas_monto = dibujar_informe(informe)
    lapiz = ImageDraw.Draw(hoja, "RGBA")
    for fila in filas_manchadas:
        x1, y1, x2, y2 = cajas_monto[fila]
        zona = (x1 + 2, y1 + 2, x2 - 2, y2 - 2)
        hoja.paste(hoja.crop(zona).filter(ImageFilter.GaussianBlur(2.5)), zona[:2])
        lapiz.ellipse((x1 - 20, y1 - 12, x2 + 20, y2 + 12), fill=(120, 85, 40, 150))
    ancho, alto = hoja.size
    sombra = Image.linear_gradient("L").rotate(90).resize(hoja.size)
    foto = Image.composite(hoja, ImageEnhance.Brightness(hoja).enhance(0.6), sombra)
    esquinas = (  # superior izq., inferior izq., inferior der., superior der.
        ancho * 0.04, alto * 0.02, ancho * 0.01, alto * 0.97,
        ancho * 0.98, alto * 0.99, ancho * 0.98, alto * 0.05,
    )
    fondo = (70, 65, 60)
    foto = foto.transform(hoja.size, Image.Transform.QUAD, esquinas, Image.Resampling.BICUBIC, fillcolor=fondo)
    foto = foto.rotate(1.5, resample=Image.Resampling.BICUBIC, fillcolor=fondo)
    foto = ImageEnhance.Contrast(foto).enhance(0.7).filter(ImageFilter.GaussianBlur(1.1))
    foto = foto.resize((int(ancho * 0.75), int(alto * 0.75)), Image.Resampling.LANCZOS)
    foto.save(ruta, "JPEG", quality=60)


def escribir_cuentas_banco(ruta: Path, cuentas: list[dict[str, str]]) -> None:
    with ruta.open("w", encoding="utf-8", newline="") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=("rut_titular", "cuenta", "codigo_banco", "estado_cuenta"),
                                  delimiter=";")
        escritor.writeheader()
        escritor.writerows(cuentas)


# --------------------------------------------------------------- generación del conjunto


def entrada_manifiesto(
    ruta_relativa: str, tipo_documento: str, proposito: str, estado_archivo: str,
    informe: Informe | None = None, motivo: str | None = None, incluir_registros: bool = True,
) -> dict[str, object]:
    entrada: dict[str, object] = {
        "archivo": ruta_relativa,
        "tipo_documento": tipo_documento,
        "proposito": proposito,
        "estado_archivo_esperado": estado_archivo,
    }
    if motivo:
        entrada["motivo"] = motivo
    if informe is not None:
        resumen: dict[str, int] = {}
        for registro in informe.registros:
            resumen[registro.estado] = resumen.get(registro.estado, 0) + 1
        entrada |= {
            "lote": informe.lote,
            "cantidad_declarada": informe.cantidad,
            "suma_declarada": informe.suma,
            "registros_en_detalle": len(informe.registros),
            "resumen_esperado": resumen,
        }
        if incluir_registros:
            entrada["registros"] = [r.para_manifiesto() for r in informe.registros]
    return entrada


def cuentas_banco(gen: Generador, registros: list[Registro]) -> list[dict[str, str]]:
    """Una cuenta activa de Banco Genérico por deudor válido (base provisoria: el cruce agregará saldos)."""
    cuentas = [{"rut_titular": r.rut, "cuenta": str(gen.rng.randint(10**9, 10**10 - 1)),
                "codigo_banco": CODIGO_BANCO_GENERICO, "estado_cuenta": "Activa"}
               for r in registros if r.estado == CARGADO]
    gen.rng.shuffle(cuentas)
    return cuentas


def generar(salida: Path, semilla: int = 2026, grande: int = 0) -> dict[str, object]:
    """Escribe todos los archivos simulados en `salida` y devuelve el manifiesto."""
    gen = Generador(semilla)
    estructurados = salida / "estructurados"
    no_estructurados = salida / "no_estructurados"
    banco = salida / "banco_generico"
    for carpeta in (estructurados, no_estructurados, banco):
        carpeta.mkdir(parents=True, exist_ok=True)

    archivos: list[dict[str, object]] = []
    deudores: list[Registro] = []

    def ruta(carpeta: Path, nombre: str) -> tuple[Path, str]:
        destino = carpeta / nombre
        return destino, destino.relative_to(salida).as_posix()

    # Formatos estructurados válidos
    destino, relativa = ruta(estructurados, "01_informe_valido.xlsx")
    informe = gen.informe(150)
    escribir_excel(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "Excel", "Informe válido en Excel, con celdas numéricas y de fecha.",
                                       "Procesado", informe))
    deudores += informe.registros
    informe_excel = relativa

    destino, relativa = ruta(estructurados, "02_informe_valido.csv")
    informe_csv = gen.informe(120)
    escribir_delimitado(destino, informe_csv)
    archivos.append(entrada_manifiesto(relativa, "CSV", "Informe válido en CSV UTF-8 separado por punto y coma.",
                                       "Procesado", informe_csv))
    deudores += informe_csv.registros

    destino, relativa = ruta(estructurados, "03_informe_latin1.txt")
    informe = gen.informe(80)
    escribir_delimitado(destino, informe, delimitador="|", encoding="latin-1")
    archivos.append(entrada_manifiesto(relativa, "TXT", "Informe válido en TXT ISO-8859-1 separado por | (RN04 "
                                       "acepta UTF-8 e ISO-8859-1).", "Procesado", informe))
    deudores += informe.registros

    destino, relativa = ruta(estructurados, "04_informe_casos_borde.xlsx")
    informe = Informe(gen.lote(), casos_de_borde(gen))
    escribir_excel(destino, informe)
    archivos.append(entrada_manifiesto(
        relativa, "Excel", "Un caso por regla de la extracción y del ETL. La suma declarada considera solo los "
        "montos numéricos del detalle.", "Procesado", informe))

    destino, relativa = ruta(estructurados, "05_informe_mes2_con_repetidos.csv")
    repetidos = [copia(r, estado=DUPLICADO, caso="Deudor ya informado en 02_informe_valido.csv",
                       motivo="Mismo RUT, expediente y fecha de ingreso que un caso ya cargado")
                 for r in gen.rng.sample(informe_csv.registros, 5)]
    informe = Informe(gen.lote(), gen.registros(55) + repetidos)
    gen.rng.shuffle(informe.registros)
    escribir_delimitado(destino, informe)
    archivos.append(entrada_manifiesto(
        relativa, "CSV", "Repite 5 deudores de 02. Son duplicados solo si llegan el mismo día que 02 (la llave del "
        "caso incluye la fecha de ingreso); un mes después la TGR los vuelve a informar y se cargan de nuevo.",
        "Procesado", informe))

    destino, relativa = ruta(estructurados, "06_informe_totales_no_cuadran.csv")
    informe = gen.informe(40)
    informe.suma_declarada = informe.suma + 125_000
    escribir_delimitado(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "CSV", "La suma de los totales no coincide con el detalle.",
                                       "Error de Extracción", informe,
                                       motivo="CU12 5b - Totales no cuadran: el informe va a revisión manual"))

    destino, relativa = ruta(estructurados, "07_informe_cantidad_no_cuadra.xlsx")
    informe = gen.informe(40)
    informe.cantidad_declarada = 43
    escribir_excel(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "Excel", "El encabezado y los totales declaran 43 deudores; el "
                                       "detalle tiene 40.", "Error de Extracción", informe,
                                       motivo="CU12 5b - Cantidad declarada distinta del detalle"))

    destino, relativa = ruta(estructurados, "08_informe_sin_totales.csv")
    informe = gen.informe(30)
    escribir_delimitado(destino, informe, con_totales=False)
    archivos.append(entrada_manifiesto(relativa, "CSV", "Sin la línea TOTALES: se carga igual; solo se verifica la "
                                       "cantidad de deudores que declara el encabezado.", "Procesado", informe))

    destino, relativa = ruta(estructurados, "09_informe_columnas_incorrectas.csv")
    informe = gen.informe(30)
    columnas = tuple(c for c in COLUMNAS if c != "expediente")
    escribir_delimitado(destino, informe, columnas=columnas, titulos={"monto_solicitado": "Importe"})
    archivos.append(entrada_manifiesto(relativa, "CSV", "Falta la columna Expediente y el monto se llama Importe.",
                                       "Error de Extracción", informe,
                                       motivo="RN04 - No se encuentra la fila de títulos del informe"))

    destino, relativa = ruta(estructurados, "10_informe_corrupto.xlsx")
    destino.write_bytes(b"PK\x03\x04" + gen.rng.randbytes(2048))
    archivos.append(entrada_manifiesto(relativa, "Excel", "Bytes aleatorios con extensión .xlsx.",
                                       "Error de Extracción", motivo="RN04 - Archivo dañado"))

    destino, relativa = ruta(estructurados, "11_informe_vacio.csv")
    informe = Informe(gen.lote(), [])
    escribir_delimitado(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "CSV", "Encabezado y totales correctos, sin deudores.",
                                       "Error de Extracción", informe, motivo="El informe no trae deudores"))

    destino, relativa = ruta(estructurados, "12_informe_valido_reenvio.xlsx")
    shutil.copyfile(salida / informe_excel, destino)
    archivos.append(entrada_manifiesto(relativa, "Excel", "Copia exacta de 01 con otro nombre.",
                                       "Rechazado (no se registra)",
                                       motivo="RN02 / CU02 4a - Archivo ya recibido (mismo contenido que 01)"))

    destino, relativa = ruta(estructurados, "13_informe_valido.docx")
    informe = gen.informe(40)
    escribir_word(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "Word", "Informe válido en Word: tabla con RUT con puntos y montos "
                                       "en pesos.", "Procesado", informe))
    deudores += informe.registros

    # Documentos no estructurados (RF12, RN05)
    destino, relativa = ruta(no_estructurados, "21_informe_pdf_texto.pdf")
    informe = gen.informe(25)
    escribir_pdf_texto(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "PDF con texto", "Tabla con texto seleccionable: se extrae sin OCR.",
                                       "Procesado", informe))
    deudores += informe.registros

    destino, relativa = ruta(no_estructurados, "22_informe_pdf_texto_totales_no_cuadran.pdf")
    informe = gen.informe(25)
    informe.suma_declarada = informe.suma - 50_000
    escribir_pdf_texto(destino, informe)
    archivos.append(entrada_manifiesto(relativa, "PDF con texto", "El total impreso no coincide con la tabla.",
                                       "Error de Extracción", informe,
                                       motivo="CU12 5b - Totales no cuadran: el informe va a revisión manual"))

    destino, relativa = ruta(no_estructurados, "23_informe_escaneado.pdf")
    informe = gen.informe(25)
    escribir_pdf_escaneado(destino, informe, gen.rng)
    archivos.append(entrada_manifiesto(
        relativa, "PDF escaneado", "Solo imagen, buena calidad. El informe se carga; los deudores con algún dato "
        "crítico (RUT, monto o expediente) leído bajo 90% de confianza quedan para revisión manual, y ninguno se "
        "acepta con un dato crítico mal leído.", "Procesado", informe))
    deudores += informe.registros

    destino, relativa = ruta(no_estructurados, "24_foto_informe.jpg")
    informe = gen.informe(25)
    manchadas = [4, 11, 19]
    for fila in manchadas:
        registro = informe.registros[fila]
        registro.estado, registro.motivo = REVISION, "RN05 - Mancha sobre el monto: confianza OCR bajo 90%"
        registro.caso = "Monto tapado por una mancha"
    escribir_fotografia(destino, informe, manchadas)
    archivos.append(entrada_manifiesto(
        relativa, "Imagen", "Foto con perspectiva, sombra y desenfoque. Los 3 montos manchados deben quedar para "
        "revisión, y ningún valor mal leído puede aceptarse sin revisión; el resto depende de la calidad real del "
        "OCR.",
        "Error de Extracción", informe,
        motivo="CU12 5b - Los montos manchados no se pueden leer, así que el detalle no cuadra con los totales: "
               "el informe completo va a revisión manual"))

    if grande:
        destino, relativa = ruta(estructurados, "90_informe_volumen.csv")
        informe = gen.informe(grande)
        escribir_delimitado(destino, informe)
        archivos.append(entrada_manifiesto(relativa, "CSV", f"Informe válido de {grande} deudores para pruebas de "
                                           "rendimiento.", "Procesado", informe, incluir_registros=False))

    # Base simulada de Banco Genérico (provisoria: el cruce con saldos y la regla de 0,5 UTM viene después)
    destino, relativa_banco = ruta(banco, "cuentas_core_bancario.csv")
    cuentas = cuentas_banco(gen, deudores)
    escribir_cuentas_banco(destino, cuentas)

    manifiesto = {
        "descripcion": "Resultados esperados de los informes de deudores simulados de TTDH Automation. Generado "
                       "por backend/scripts/generar_datos_simulados.py; no editar a mano.",
        "semilla": semilla,
        "informe": {
            "encabezado": [TITULO_INFORME, AVISO, "Lote: <lote>  Fecha de generación: <dd-mm-aaaa>  Registros: <n>"],
            "columnas": {c: TITULOS[c] for c in COLUMNAS},
            "totales": "TOTALES  Registros: <n>  Monto total: $ <suma>",
            "estados_esperados": [CARGADO, REVISION, RECHAZADO, DUPLICADO],
        },
        "archivos": archivos,
        "banco_generico": {"archivo": relativa_banco, "cuentas": len(cuentas)},
    }
    (salida / "manifiesto.json").write_text(json.dumps(manifiesto, ensure_ascii=False, indent=2) + "\n",
                                            encoding="utf-8")
    return manifiesto


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Genera los informes de deudores simulados de TTDH Automation.")
    parser.add_argument("--salida", type=Path, default=CARPETA_SALIDA,
                        help=f"carpeta de destino (por defecto {CARPETA_SALIDA})")
    parser.add_argument("--semilla", type=int, default=2026, help="semilla aleatoria (por defecto 2026)")
    parser.add_argument("--grande", type=int, default=0, metavar="N",
                        help="agrega un informe CSV válido de N deudores para pruebas de volumen")
    args = parser.parse_args(argv)
    if args.grande < 0:
        parser.error("--grande debe ser 0 o mayor")

    salida = args.salida.resolve()
    print(f"Generando datos simulados en {salida}")
    manifiesto = generar(salida, args.semilla, args.grande)
    for archivo in manifiesto["archivos"]:
        resumen = archivo.get("resumen_esperado")
        detalle = ", ".join(f"{estado}: {n}" for estado, n in resumen.items()) if resumen else "-"
        print(f"  {archivo['archivo']:<55} {archivo['estado_archivo_esperado']:<27} {detalle}")
    banco = manifiesto["banco_generico"]
    print(f"  {banco['archivo']:<55} {banco['cuentas']} cuentas")
    print(f"Listo: {len(manifiesto['archivos'])} informes y manifiesto.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
