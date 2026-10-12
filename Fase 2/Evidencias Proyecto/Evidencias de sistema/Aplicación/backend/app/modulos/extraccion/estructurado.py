"""Lectura de informes de deudores en formatos estructurados: Excel, CSV, TXT y Word (RF01, RF02, RN04).

No hay OCR: los valores vienen tal como están en el archivo, con confianza 100. Se busca la fila de
títulos (puede haber líneas de encabezado antes), se leen los deudores hasta la línea TOTALES y con el
resto del texto se leen el encabezado y las sumatorias de control, igual que en un PDF.
"""
import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path

from .documento import columna_desde_titulo, leer_encabezado_y_totales
from .normalizacion import limpiar, normalizar
from .plantilla import COLUMNAS_PLANTILLA, RegistroExtraido, ResultadoExtraccion

TIPOS = {".xlsx": "Excel", ".csv": "CSV", ".txt": "TXT", ".docx": "Word"}
_DELIMITADORES = ";|,\t"


class ArchivoIlegible(Exception):
    """El archivo está dañado o no corresponde a su extensión."""


def _texto_celda(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d-%m-%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _filas_excel(ruta: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as error:
        raise ArchivoIlegible("El Excel está dañado o no es un .xlsx válido.") from error
    try:
        return [[_texto_celda(c) for c in fila] for fila in libro.active.iter_rows(values_only=True)]
    finally:
        libro.close()


def _filas_delimitadas(ruta: Path) -> list[list[str]]:
    """CSV o TXT en UTF-8 o ISO-8859-1 (RN04), separado por ; | , o tabulador."""
    datos = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    # El separador es el que divide más líneas en al menos tantas partes como columnas tiene el informe.
    # (csv.Sniffer se confunde con las líneas de encabezado y totales, que tienen pocas celdas.)
    lineas = texto.splitlines()[:50]
    delimitador = max(_DELIMITADORES,
                      key=lambda d: sum(linea.count(d) >= len(COLUMNAS_PLANTILLA) - 1 for linea in lineas))
    return list(csv.reader(io.StringIO(texto), delimiter=delimitador))


def _filas_word(ruta: Path) -> list[list[str]]:
    """Párrafos (encabezado y totales) y filas de las tablas, en el orden del documento."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        documento = Document(str(ruta))
    except (zipfile.BadZipFile, KeyError, ValueError) as error:
        raise ArchivoIlegible("El Word está dañado o no es un .docx válido.") from error
    filas: list[list[str]] = []
    for bloque in documento.iter_inner_content():
        if isinstance(bloque, Paragraph):
            filas.append([bloque.text])
        elif isinstance(bloque, Table):
            filas.extend([celda.text for celda in fila.cells] for fila in bloque.rows)
    return filas


def leer_filas(ruta: Path) -> list[list[str]]:
    extension = ruta.suffix.lower()
    if extension == ".xlsx":
        return _filas_excel(ruta)
    if extension in (".csv", ".txt"):
        return _filas_delimitadas(ruta)
    if extension == ".docx":
        return _filas_word(ruta)
    raise ArchivoIlegible(f"Formato {extension!r} no soportado.")


def extraer_estructurado(ruta: Path) -> ResultadoExtraccion:
    resultado = ResultadoExtraccion(tipo_documento=TIPOS.get(ruta.suffix.lower(), ruta.suffix))
    try:
        filas = leer_filas(ruta)
    except ArchivoIlegible as error:
        resultado.errores.append(str(error))
        return resultado

    columnas: list[str | None] | None = None
    fuera_de_tabla: list[str] = []
    en_totales = False
    for fila in filas:
        texto = limpiar(" ".join(fila))
        if not texto:
            continue
        posibles = [columna_desde_titulo(celda) for celda in fila]
        if columnas is None and set(COLUMNAS_PLANTILLA) <= set(posibles):
            columnas = posibles
            continue
        if columnas is None or en_totales or texto.upper().startswith("TOTALES"):
            en_totales = en_totales or (columnas is not None and texto.upper().startswith("TOTALES"))
            fuera_de_tabla.append(texto)
            continue
        valores = {col: normalizar(col, celda) for col, celda in zip(columnas, fila) if col}
        resultado.registros.append(RegistroExtraido(valores, {col: 100.0 for col in COLUMNAS_PLANTILLA}))

    leer_encabezado_y_totales(resultado, "\n".join(fuera_de_tabla))
    if columnas is None:
        faltan = ", ".join(COLUMNAS_PLANTILLA)
        resultado.errores.append(
            f"No se encontró la fila de títulos del informe (se esperan las columnas: {faltan}).")
        return resultado
    if not resultado.registros:
        resultado.errores.append("El informe no trae deudores.")
        return resultado
    resultado.verificar_totales()
    return resultado
