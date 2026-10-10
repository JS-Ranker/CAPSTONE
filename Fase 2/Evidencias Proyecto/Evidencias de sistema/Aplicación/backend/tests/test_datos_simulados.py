"""Pruebas del generador de datos simulados: los archivos deben cumplir lo que dice el manifiesto."""
import csv
import json

import pdfplumber
import pytest
from openpyxl import load_workbook

from scripts.generar_datos_simulados import COLUMNAS, digito_verificador, generar

CAMPOS_OBLIGATORIOS = ("monto", "cuenta_destino", "codigo_banco", "tipo_movimiento")  # RN03


@pytest.fixture(scope="module")
def salida(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("simulados")
    generar(carpeta, semilla=2026)
    return carpeta


@pytest.fixture(scope="module")
def manifiesto(salida):
    return json.loads((salida / "manifiesto.json").read_text(encoding="utf-8"))


def entrada(manifiesto, nombre):
    return next(a for a in manifiesto["archivos"] if a["archivo"].endswith(nombre))


def leer_delimitado(ruta, delimitador=";", encoding="utf-8"):
    with ruta.open(encoding=encoding, newline="") as archivo:
        filas = list(csv.reader(archivo, delimiter=delimitador))
    return filas[0], filas[1], filas[2:-1], filas[-1]


def rut_valido(rut):
    numero, _, dv = rut.partition("-")
    return numero.isdigit() and dv != "" and digito_verificador(int(numero)) == dv.upper()


def clave_rn02(registro):
    return registro["id_transaccion"], registro["rut_deudor"], registro["monto"], registro["fecha_origen"]


@pytest.mark.parametrize("numero, dv", [(11111111, "1"), (12345678, "5"), (22222222, "2"), (10000013, "K")])
def test_digito_verificador_modulo_11(numero, dv):
    assert digito_verificador(numero) == dv


def test_todos_los_archivos_del_manifiesto_existen(salida, manifiesto):
    for archivo in manifiesto["archivos"]:
        assert (salida / archivo["archivo"]).is_file(), archivo["archivo"]
    assert (salida / manifiesto["cruce_banco_generico"]["archivo"]).is_file()


def test_ruts_validos_fuera_del_rango_de_personas_y_empresas(manifiesto):
    for archivo in manifiesto["archivos"]:
        for registro in archivo.get("registros", []):
            if registro["estado_esperado"] == "Válido":
                assert rut_valido(registro["rut_deudor"]), registro
                assert 40_000_000 <= int(registro["rut_deudor"].split("-")[0]) <= 49_999_999


def test_csv_valido_cuadra_con_encabezado_y_pie(salida):
    encabezado, columnas, detalle, pie = leer_delimitado(salida / "estructurados/02_nomina_valida.csv")
    assert tuple(columnas) == COLUMNAS
    assert int(encabezado[4]) == len(detalle) == int(pie[1]) == 120
    assert sum(int(fila[4]) for fila in detalle) == int(pie[2])


def test_txt_en_latin1_separado_por_barra(salida):
    ruta = salida / "estructurados/03_nomina_latin1.txt"
    assert "Depósito".encode("latin-1") in ruta.read_bytes()
    _, columnas, detalle, pie = leer_delimitado(ruta, delimitador="|", encoding="latin-1")
    assert tuple(columnas) == COLUMNAS
    assert sum(int(fila[4]) for fila in detalle) == int(pie[2])


def test_excel_valido_con_celdas_tipadas(salida):
    hoja = load_workbook(salida / "estructurados/01_nomina_valida.xlsx").active
    filas = list(hoja.iter_rows(values_only=True))
    encabezado, columnas, detalle, pie = filas[0], filas[1], filas[2:-1], filas[-1]
    assert tuple(columnas) == COLUMNAS
    assert int(encabezado[4]) == len(detalle) == int(pie[1])
    assert all(isinstance(fila[4], int) for fila in detalle)
    assert sum(fila[4] for fila in detalle) == int(pie[2])


def test_errores_de_estructura(salida):
    *_, detalle, pie = leer_delimitado(salida / "estructurados/06_nomina_totales_no_cuadran.csv")
    assert sum(int(fila[4]) for fila in detalle) != int(pie[2])

    *_, ultima = leer_delimitado(salida / "estructurados/08_nomina_sin_pie.csv")
    assert ultima[0] != "TOTALES"

    _, columnas, _, _ = leer_delimitado(salida / "estructurados/09_nomina_columnas_incorrectas.csv")
    assert "cuenta_destino" not in columnas and "importe" in columnas

    with pytest.raises(Exception):
        load_workbook(salida / "estructurados/10_nomina_corrupta.xlsx")


def test_reenvio_es_identico_al_original(salida):
    original = (salida / "estructurados/01_nomina_valida.xlsx").read_bytes()
    assert (salida / "estructurados/12_nomina_valida_reenvio.xlsx").read_bytes() == original


def test_casos_de_borde_coinciden_con_las_reglas(manifiesto):
    registros = entrada(manifiesto, "04_nomina_casos_borde.xlsx")["registros"]
    vistos = set()
    for registro in registros:
        estado, motivo = registro["estado_esperado"], registro.get("motivo", "")
        if estado == "Válido":
            assert rut_valido(registro["rut_deudor"]), registro
            assert all(registro[campo] not in (None, "") for campo in CAMPOS_OBLIGATORIOS), registro
        if motivo.startswith("RN01"):
            assert not rut_valido(registro["rut_deudor"]), registro
        if motivo.startswith("RN03"):
            assert any(registro[campo] in (None, "") for campo in CAMPOS_OBLIGATORIOS), registro
        if estado == "Duplicado":
            assert clave_rn02(registro) in vistos, registro
        vistos.add(clave_rn02(registro))
    assert {r["estado_esperado"] for r in registros} == {"Válido", "Inválido", "Duplicado"}


def test_lote2_repite_transacciones_del_csv_valido(manifiesto):
    claves_02 = {clave_rn02(r) for r in entrada(manifiesto, "02_nomina_valida.csv")["registros"]}
    duplicados = [r for r in entrada(manifiesto, "05_nomina_lote2_con_duplicados.csv")["registros"]
                  if r["estado_esperado"] == "Duplicado"]
    assert len(duplicados) == 5
    assert all(clave_rn02(r) in claves_02 for r in duplicados)


def test_pdf_con_texto_contiene_la_tabla(salida, manifiesto):
    with pdfplumber.open(salida / "no_estructurados/21_nomina_pdf_texto.pdf") as pdf:
        tabla = pdf.pages[0].extract_table()
    esperados = entrada(manifiesto, "21_nomina_pdf_texto.pdf")["registros"]
    assert [fila[0] for fila in tabla[1:]] == [r["id_transaccion"] for r in esperados]


def test_pdf_escaneado_no_tiene_texto_seleccionable(salida):
    with pdfplumber.open(salida / "no_estructurados/23_nomina_escaneada.pdf") as pdf:
        pagina = pdf.pages[0]
        assert not (pagina.extract_text() or "").strip()
        assert pagina.images


def test_cruce_banco_generico(salida, manifiesto):
    cruce = manifiesto["cruce_banco_generico"]
    with (salida / cruce["archivo"]).open(encoding="utf-8", newline="") as archivo:
        cuentas = {fila["cuenta"]: fila for fila in csv.DictReader(archivo, delimiter=";")}
    assert len(cuentas) == cruce["cuentas"]
    for caso in cruce["casos"]:
        cuenta = cuentas.get(caso["cuenta_destino"])
        if caso["caso"] == "Cuenta de destino cerrada":
            assert cuenta["estado_cuenta"] == "Cerrada"
        elif caso["caso"] == "Cuenta de destino inexistente en Banco Genérico":
            assert cuenta is None
        else:
            assert cuenta["rut_titular"] != caso["rut_deudor"]


def test_misma_semilla_produce_los_mismos_datos(salida, tmp_path):
    generar(tmp_path, semilla=2026)
    for nombre in ("manifiesto.json", "estructurados/02_nomina_valida.csv", "banco_generico/cuentas_core_bancario.csv"):
        assert (tmp_path / nombre).read_bytes() == (salida / nombre).read_bytes(), nombre


def test_nomina_de_volumen(tmp_path):
    manifiesto = generar(tmp_path, semilla=1, grande=500)
    assert "registros" not in entrada(manifiesto, "90_nomina_volumen.csv")
    *_, detalle, pie = leer_delimitado(tmp_path / "estructurados/90_nomina_volumen.csv")
    assert len(detalle) == int(pie[1]) == 500
